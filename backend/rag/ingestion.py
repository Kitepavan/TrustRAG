import os
import hashlib
import uuid
from pathlib import Path
from datetime import datetime, timezone
from typing import Optional

import fitz  # PyMuPDF
import docx  # python-docx

from backend.rag.chunking import chunk_document
from backend.security.hashing import compute_sha256_file
from backend.security.signature import verify_signature
from backend.security.trust_engine import trust_engine
from backend.security.provenance import provenance_tracker, ProvenanceRecord


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE_MB = 50

# Document classification levels. Must match the RBAC clearance tags in auth.py:
# a chunk's access_level is what Defense Filter 2 enforces at retrieval time.
ALLOWED_ACCESS_LEVELS = {"PUBLIC", "INTERNAL", "HR_CONFIDENTIAL", "IT_SEC_CONFIDENTIAL", "RESTRICTED"}


def validate_file(file_path: str, filename: str) -> dict:
    """Validate uploaded file for type, size, and readability."""
    errors = []
    path = Path(file_path)

    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        errors.append(f"Unsupported file type: {ext}. Allowed: {ALLOWED_EXTENSIONS}")

    size_mb = os.path.getsize(file_path) / (1024 * 1024)
    if size_mb > MAX_FILE_SIZE_MB:
        errors.append(f"File too large: {size_mb:.1f}MB. Max: {MAX_FILE_SIZE_MB}MB")

    if not os.access(file_path, os.R_OK):
        errors.append("File is not readable")

    return {"valid": len(errors) == 0, "errors": errors}


def calculate_sha256(file_path: str) -> str:
    """Calculate SHA-256 hash of a file."""
    return compute_sha256_file(file_path)


def extract_text_pdf(file_path: str) -> dict:
    """Extract text from PDF using PyMuPDF."""
    doc = fitz.open(file_path)
    pages = []
    full_text = []

    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text = page.get_text()
        pages.append({
            "page_number": page_num + 1,
            "text": text,
            "char_count": len(text),
        })
        full_text.append(text)

    doc.close()

    return {
        "pages": pages,
        "total_pages": len(pages),
        "full_text": "\n".join(full_text),
        "total_chars": sum(p["char_count"] for p in pages),
    }


def extract_text_docx(file_path: str) -> dict:
    """Extract text from DOCX using python-docx."""
    doc = docx.Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                tables.append(" | ".join(cells))

    full_text = "\n".join(paragraphs)
    if tables:
        full_text = f"{full_text}\n\n" + "\n".join(tables)

    return {
        "pages": [{"page_number": 1, "text": full_text, "char_count": len(full_text)}],
        "total_pages": 1,
        "full_text": full_text,
        "total_chars": len(full_text),
    }


def extract_text(file_path: str) -> dict:
    """Extract text based on file type."""
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        return extract_text_pdf(file_path)
    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
        return {
            "pages": [{"page_number": 1, "text": text, "char_count": len(text)}],
            "total_pages": 1,
            "full_text": text,
            "total_chars": len(text),
        }
    elif ext == ".docx":
        return extract_text_docx(file_path)
    else:
        raise ValueError(f"Text extraction not implemented for {ext}")


def ingest_document(
    file_path: str,
    filename: str,
    uploaded_by: str = "system_admin",
    signature_b64: Optional[str] = None,
    chunk_size: int = 512,
    chunk_overlap: int = 64,
    access_level: str = "INTERNAL",
) -> dict:
    """
    Full TrustRAG Ingestion Pipeline:
    Validate → Hash → Extract Text → Verify Signature → Trust Engine Evaluation → Record Provenance → Chunk
    """
    validation = validate_file(file_path, filename)
    if not validation["valid"]:
        return {"status": "error", "errors": validation["errors"]}

    access_level = (access_level or "INTERNAL").strip().upper()
    if access_level not in ALLOWED_ACCESS_LEVELS:
        return {
            "status": "error",
            "errors": [f"Invalid access_level: {access_level}. Allowed: {sorted(ALLOWED_ACCESS_LEVELS)}"],
        }

    doc_id = f"DOC-{uuid.uuid4().hex.upper()}"
    sha256_hash = calculate_sha256(file_path)

    is_signed = signature_b64 is not None and len(signature_b64) > 0
    raw_bytes = b""
    # Read binary content for signature verification
    if is_signed:
        with open(file_path, "rb") as f:
            raw_bytes = f.read()

    extraction = extract_text(file_path)

    if extraction["total_chars"] > 2_000_000:
        return {"status": "error", "errors": ["Extracted text exceeds 2 million characters"]}

    signature_valid = verify_signature(raw_bytes, signature_b64) if is_signed else False

    # Trust Engine Evaluation
    trust_eval = trust_engine.evaluate_document_trust(
        file_bytes=raw_bytes,
        text_content=extraction["full_text"],
        is_signed=is_signed,
        signature_valid=signature_valid,
        uploader=uploaded_by,
    )

    trust_status = trust_eval["trust_category"]
    trust_score = trust_eval["trust_score"]

    # Record provenance log
    prov_record = ProvenanceRecord(
        document_id=doc_id,
        filename=filename,
        sha256_hash=sha256_hash,
        signature_valid=signature_valid,
        uploader=uploaded_by,
        trust_status=trust_status,
        notes="; ".join(trust_eval["reasons"]),
    )
    provenance_tracker.record_provenance(prov_record)

    extraction["document_id"] = doc_id
    chunks = chunk_document(extraction, chunk_size=chunk_size, chunk_overlap=chunk_overlap)

    return {
        "status": "processed",
        "document_id": doc_id,
        "filename": filename,
        "sha256": sha256_hash,
        "is_signed": is_signed,
        "signature_valid": signature_valid,
        "trust_status": trust_status,
        "trust_score": trust_score,
        "policy_decision": trust_eval["policy_decision"],
        "reasons": trust_eval["reasons"],
        "uploaded_by": uploaded_by,
        "uploaded_at": prov_record.upload_timestamp,
        "access_level": access_level,
        "total_pages": extraction["total_pages"],
        "total_chars": extraction["total_chars"],
        "total_chunks": len(chunks),
        "chunks": [
            {
                "chunk_id": c.chunk_id,
                "document_id": doc_id,
                "text": c.text,
                "page_number": c.page_number,
                "char_count": c.char_count,
                "sha256": sha256_hash,
                "signature_valid": signature_valid,
                "trust_status": trust_status,
                "trust_score": trust_score,
                "access_level": access_level,
            }
            for c in chunks
        ],
    }
